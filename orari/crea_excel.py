"""Crea il file Excel degli orari a partire da piano.json."""
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from genera_orari import GIORNI, PERSONE, REPARTO

piano = json.load(open("piano.json"))
wb = Workbook()
ws = wb.active
ws.title = "Orari"
F = "Arial"
thin = Side(style="thin", color="808080")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
GRIGIO = PatternFill("solid", fgColor="D9D9D9")
RIP = PatternFill("solid", fgColor="FCE4D6")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws["A1"] = "DAL"; ws["B1"] = "12/10/2026"; ws["D1"] = "AL"; ws["E1"] = "18/10/2026"
for c in ("A1", "B1", "D1", "E1"):
    ws[c].font = Font(name=F, bold=True, size=12)

# intestazioni (righe 3-5)
ws.merge_cells("A3:A5"); ws["A3"] = "NOMINATIVO / TURNO"
ws.merge_cells("B3:B5"); ws["B3"] = "RUOLO"
for d, g in enumerate(GIORNI):
    c = 3 + 4 * d
    ws.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + 3)
    ws.cell(3, c, g)
    ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
    ws.cell(4, c, "Mattino")
    ws.merge_cells(start_row=4, start_column=c + 2, end_row=4, end_column=c + 3)
    ws.cell(4, c + 2, "Pomeriggio")
    for k, t in enumerate("EUEU"):
        ws.cell(5, c + k, t)
TOT = 3 + 4 * 7
ws.merge_cells(start_row=3, start_column=TOT, end_row=5, end_column=TOT); ws.cell(3, TOT, "TOTALE ORE")
ws.merge_cells(start_row=3, start_column=TOT + 1, end_row=5, end_column=TOT + 1); ws.cell(3, TOT + 1, "ORE CONTRATTO")
for r in range(3, 6):
    for c in range(1, TOT + 2):
        cell = ws.cell(r, c)
        cell.font = Font(name=F, bold=True, size=9); cell.alignment = CEN; cell.border = B; cell.fill = GRIGIO

r0 = 6
for n, (nome, ruolo, ore, chiavi, _) in enumerate(PERSONE):
    r = r0 + n
    ws.cell(r, 1, nome + (" 🔑" if chiavi else "")).font = Font(name=F, bold=True, size=9)
    ws.cell(r, 2, ruolo).font = Font(name=F, size=9)
    for d in range(7):
        c = 3 + 4 * d
        t = piano[nome].get(str(d))
        if t is None:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 3)
            ws.cell(r, c, "RIPOSO")
            for k in range(4):
                ws.cell(r, c + k).fill = RIP
            continue
        if len(t) == 2:
            slots = {0: t[0], 2: t[1]}
        else:
            slots = {0 if t[0][0] < 13 else 2: t[0]}
        for k, (s, e) in slots.items():
            ws.cell(r, c + k, s / 24).number_format = "hh:mm"
            ws.cell(r, c + k + 1, e / 24).number_format = "hh:mm"
    parts = []
    for d in range(7):
        a = [L(3 + 4 * d + k) + str(r) for k in range(4)]
        parts.append(f"N({a[1]})-N({a[0]})+N({a[3]})-N({a[2]})")
    ws.cell(r, TOT, "=ROUND(24*(" + "+".join(parts) + "),2)")
    ws.cell(r, TOT + 1, "44-45" if nome == "DE SANTIS MIRKO" else ore)
    for c in range(1, TOT + 2):
        cell = ws.cell(r, c)
        cell.border = B
        if c > 2:
            cell.alignment = CEN
            cell.font = Font(name=F, size=9, bold=(c == TOT))
r_last = r0 + len(PERSONE) - 1

note = [
    "🔑 = ha le chiavi (apre/chiude). Ogni giorno un tesserato apre alle 9:00 e uno chiude alle 21:00.",
    "Riposi spostati rispetto alla settimana 5-11 ottobre: mai lo stesso giorno, distanza tra un riposo e l'altro da 2 a 5 giorni.",
    "Orari delle casse (Benedetta, Martina, Debora, Giorgia) presi dal foglio del negozio; Luigi Diana lun-ven 9-13 / 14-18.",
    "Sabato giornata intera per tutto il reparto; il terzo addetto di reparto messo di preferenza al pomeriggio.",
    "Il foglio 'Copertura' conta ora per ora quante persone ci sono (si aggiorna da solo se modifichi gli orari).",
]
for i, t in enumerate(note):
    ws.cell(r_last + 2 + i, 1, t).font = Font(name=F, italic=True, size=9)

ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 15
for c in range(3, TOT):
    ws.column_dimensions[L(c)].width = 6.2
ws.column_dimensions[L(TOT)].width = 9
ws.column_dimensions[L(TOT + 1)].width = 10
ws.freeze_panes = "C6"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# --- Foglio copertura ---------------------------------------------------
cv = wb.create_sheet("Copertura")
cv["A1"] = "Persone presenti ora per ora (calcolate dal foglio Orari)"
cv["A1"].font = Font(name=F, bold=True, size=12)
ruoli = f"Orari!$B${r0}:$B${r_last}"
nomi = f"Orari!$A${r0}:$A${r_last}"
blocchi = [("ADDETTI REPARTO (vendite + vicedirettore) — obiettivo almeno 2",
            f'((({ruoli}="Addetto vendite")+({ruoli}="Vicedirettore"))>0)', 2),
           ("CASSIERE — obiettivo almeno 1", f'({ruoli}="Cassiera")', 1),
           ("PERSONE CON LE CHIAVI — obiettivo almeno 1", f'(ISNUMBER(SEARCH("🔑",{nomi})))', 1)]
row = 3
ROSSO = PatternFill("solid", fgColor="F8CBAD")
from openpyxl.formatting.rule import CellIsRule
for titolo, filtro, minimo in blocchi:
    cv.cell(row, 1, titolo).font = Font(name=F, bold=True)
    row += 1
    cv.cell(row, 1, "Giorno").font = Font(name=F, bold=True)
    for j, h in enumerate(range(9, 21)):
        cell = cv.cell(row, 2 + j, f"{h}-{h+1}")
        cell.font = Font(name=F, bold=True, size=9); cell.alignment = CEN; cell.fill = GRIGIO; cell.border = B
    row += 1
    start = row
    for d, g in enumerate(GIORNI):
        cv.cell(row, 1, g).font = Font(name=F, size=9)
        c = 3 + 4 * d
        rng = [f"Orari!${L(c + k)}${r0}:${L(c + k)}${r_last}" for k in range(4)]
        for j, h in enumerate(range(9, 21)):
            t = f"TIME({h},30,0)"
            f = (f"=SUMPRODUCT({filtro}*((({rng[0]}<={t})*({rng[1]}>{t}))"
                 f"+(({rng[2]}<={t})*({rng[3]}>{t}))))")
            cell = cv.cell(row, 2 + j, f)
            cell.alignment = CEN; cell.border = B; cell.font = Font(name=F, size=9)
        row += 1
    cv.conditional_formatting.add(f"B{start}:M{row-1}",
                                  CellIsRule(operator="lessThan", formula=[str(minimo)], fill=ROSSO))
    row += 1
cv.column_dimensions["A"].width = 26
for j in range(12):
    cv.column_dimensions[L(2 + j)].width = 7
wb.save("Orari_12-18_ottobre_2026.xlsx")
print("ok")
